"""html2xlsx 引擎单元测试：用合成 HTML 校验结构转写与样式映射的关键行为。"""

from __future__ import annotations

from pathlib import Path

import pytest
from openpyxl import load_workbook

from baibao.render.html2xlsx import (
    BorderStyle,
    CellStyle,
    Html2XlsxProfile,
    ImageAnchor,
    convert_html_to_xlsx,
)

# 1x1 透明 PNG（最小合法 PNG，仅用于锚定计数，不校验像素）
_TINY_PNG = bytes.fromhex(
    "89504e470d0a1a0a0000000d49484452000000010000000108060000001f15c489"
    "0000000d49444154789c626001000000ffff03000006000557bfabd40000000049454e44ae426082"
)


def _convert(tmp_path: Path, html: str, profile: Html2XlsxProfile | None = None, **kwargs):
    out = tmp_path / "out.xlsx"
    result = convert_html_to_xlsx(html, out, profile=profile, **kwargs)
    return result, load_workbook(out)


def test_table_structure_and_merges(tmp_path: Path) -> None:
    """colspan/rowspan 展开为网格并生成合并区，行列数正确。"""
    html = """
    <div class="doc">
      <table class="goods-table">
        <tr>
          <th style="width:50%">名称</th><th style="width:25%">数量</th><th style="width:25%">金额</th>
        </tr>
        <tr><td rowspan="2">款A</td><td>10</td><td>100</td></tr>
        <tr><td colspan="2">共计：110</td></tr>
      </table>
    </div>
    """
    result, wb = _convert(tmp_path, html)
    ws = wb.active
    assert result.cols == 3
    assert ws.cell(row=1, column=1).value == "名称"
    assert ws.cell(row=2, column=1).value == "款A"
    assert ws.cell(row=3, column=2).value == "共计：110"
    merged = {str(r) for r in ws.merged_cells.ranges}
    assert "A2:A3" in merged
    assert "B3:C3" in merged


def test_style_mapping_font_align_fill_border(tmp_path: Path) -> None:
    """class 映射与行内样式落格：粗体/右对齐/灰底/全线框。"""
    html = """
    <table class="t">
      <tr><th>名称</th><th>金额</th></tr>
      <tr><td class="lft">洗标</td><td class="num" style="color:#FF0000;background:#F3F3F3">¥10.00</td></tr>
    </table>
    """
    profile = Html2XlsxProfile(
        class_styles={
            ".num": CellStyle(halign="right"),
            ".lft": CellStyle(halign="left"),
        }
    )
    result, wb = _convert(tmp_path, html, profile)
    ws = wb.active
    assert result.warnings == []
    num_cell = ws.cell(row=2, column=2)
    assert not num_cell.font.bold  # 非表头不加粗
    assert num_cell.alignment.horizontal == "right"
    assert num_cell.fill.fgColor.rgb == "FFF3F3F3"
    assert num_cell.font.color.rgb == "FFFF0000"
    th = ws.cell(row=1, column=1)
    assert th.font.bold is True
    assert th.alignment.horizontal == "center"
    # 全线框：四边都有细边框
    assert ws.cell(row=2, column=1).border.top.style == "thin"
    assert ws.cell(row=2, column=1).border.bottom.style == "thin"
    # px → pt
    assert ws.cell(row=2, column=1).font.size == pytest.approx(9.0)


def test_border_none_and_block_bottom_border(tmp_path: Path) -> None:
    """行内 border:none 覆盖全线框；段落 class 的 border-bottom 落到合并行。"""
    html = """
    <div class="doc">
      <table class="t">
        <tr><td style="border-bottom:none">无底线</td><td>普通</td></tr>
      </table>
      <div class="clause-item">第一条：测试条款</div>
    </div>
    """
    profile = Html2XlsxProfile(
        container_class="doc",
        class_styles={".clause-item": CellStyle(border_bottom=BorderStyle(style="thin"))},
    )
    _, wb = _convert(tmp_path, html, profile)
    ws = wb.active
    assert ws.cell(row=1, column=1).border.bottom.style is None
    assert ws.cell(row=1, column=2).border.bottom.style == "thin"
    # 条款行：合并整行宽 + 底边框
    assert ws.cell(row=3, column=1).border.bottom.style == "thin"
    assert any(str(r).startswith("A3:") for r in ws.merged_cells.ranges)


def test_column_widths_from_percent(tmp_path: Path) -> None:
    """首行 width% 按总预算分摊列宽，未标注列均摊剩余。"""
    html = """
    <table class="t">
      <tr><th style="width:75%">甲</th><th>乙</th></tr>
      <tr><td>a</td><td>b</td></tr>
    </table>
    """
    profile = Html2XlsxProfile(total_width_chars=100.0)
    _, wb = _convert(tmp_path, html, profile)
    ws = wb.active
    w1 = ws.column_dimensions["A"].width
    w2 = ws.column_dimensions["B"].width
    assert w1 == pytest.approx(75.0)
    assert w2 == pytest.approx(25.0)
    assert w1 > w2


def test_column_layout_two_bands(tmp_path: Path) -> None:
    """多栏容器（flex 近似）：左右两块并排写入不同列段。"""
    html = """
    <div class="doc">
      <div class="head">
        <div class="left"><p>需方：甲公司</p></div>
        <div class="right"><p>合同编号：X1</p></div>
      </div>
    </div>
    """
    profile = Html2XlsxProfile(
        container_class="doc",
        column_layouts={"head": (1.0, 1.0)},
    )
    result, wb = _convert(tmp_path, html, profile)
    ws = wb.active
    texts = {(c.row, c.column): c.value for row in ws.iter_rows(min_row=1, max_row=3) for c in row}
    values = [v for v in texts.values() if v]
    assert any("需方" in v for v in values)
    assert any("合同编号" in v for v in values)
    assert result.warnings == []


def test_image_anchor_with_profile_rule(tmp_path: Path) -> None:
    """表格内图片经 image_rule 锚定到指定单元格。"""
    html = """
    <table class="sign">
      <tr><td>供方：<img class="seal" src="seal.png"></td><td>需方</td></tr>
    </table>
    """
    def seal_rule(img, containing):
        if "seal" in (img.get("class") or []):
            return ImageAnchor(row=0, col=0, width_px=120, height_px=120)
        return None

    profile = Html2XlsxProfile(image_rule=seal_rule)
    result, wb = _convert(tmp_path, html, profile, images={"seal.png": _TINY_PNG})
    ws = wb.active
    assert result.images == 1
    assert len(ws._images) == 1


def test_unknown_class_blocks_degrade_with_warning(tmp_path: Path) -> None:
    """未知 class 的文本块降级为普通文本行，不报错但给 warnings。"""
    html = "<div class='mystery'>神秘块</div><p>普通段落</p>"
    profile = Html2XlsxProfile(class_styles={})
    result, wb = _convert(tmp_path, html, profile)
    ws = wb.active
    assert ws.cell(row=1, column=1).value == "神秘块"
    assert ws.cell(row=2, column=1).value == "普通段落"
    assert isinstance(result.warnings, list)


def test_position_rule_via_table_style_rule(tmp_path: Path) -> None:
    """table_style_rule 位置回调：第 2 行第 1 列去底线（签署栏 nth-child 语义）。"""
    html = """
    <table class="sign-table">
      <tr><td>供方</td><td>需方</td><td>鉴证</td></tr>
      <tr><td>占位</td><td>名称</td><td>意见</td></tr>
    </table>
    """
    def rule(table_cls, row_idx, col_idx, style, total_rows):
        if table_cls == "sign-table" and row_idx == 1 and col_idx == 0:
            return style.merge(CellStyle(border_bottom=BorderStyle(style=None)))
        return style

    profile = Html2XlsxProfile(table_style_rule=rule)
    _, wb = _convert(tmp_path, html, profile)
    ws = wb.active
    assert ws.cell(row=2, column=1).border.bottom.style is None
    assert ws.cell(row=2, column=2).border.bottom.style == "thin"
